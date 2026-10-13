import re
import numpy as np
from collections import defaultdict
from openpyxl import load_workbook

# Default target metric labels in Mouseframe Excel tables
DEFAULT_RAW_STRINGS = [
    "Stride length left front average in cm:",
    "Stride length right front average in cm:",
    "Stride length left hind average in cm:",
    "Stride length right hind average in cm:",
    "Overlap left average in cm:",
    "Overlap Right average in cm:",
    "Stride Width Front(L) average in cm:",
    "Stride Width Front(R) average in cm:",
    "Stride Width Hind(L) average in cm:",
    "Stride Width Hind(R) average in cm:"
]

# Rename mapping for unified metric labels
DEFAULT_RENAME_MAP = {
    "Stride Width Front(L) average in cm:": "Stride Width Front average in cm:",
    "Stride Width Front(R) average in cm:": "Stride Width Front average in cm:",
    "Stride Width Hind(L) average in cm:": "Stride Width Hind average in cm:",
    "Stride Width Hind(R) average in cm:": "Stride Width Hind average in cm:"
}


def find_table_data(workbook_path, threshold=500.0, raw_strings=None, rename_map=None):
    """
    Extracts table data from a raw Mouseframe Excel workbook.
    """
    if raw_strings is None:
        raw_strings = DEFAULT_RAW_STRINGS
    if rename_map is None:
        rename_map = DEFAULT_RENAME_MAP

    wb = load_workbook(workbook_path, data_only=True)
    result_data = []

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        current_sheet_tag = sheet_name

        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.strip().startswith("Table ID"):
                    raw_table_id = cell.value.strip().replace("Table ID:", "").strip()

                    table_values = defaultdict(list)
                    table_values["Table ID"].append(raw_table_id)
                    table_values["Sheet_Source"].append(current_sheet_tag)

                    col_index = cell.column
                    row_index = cell.row

                    for r in range(row_index + 1, ws.max_row + 1):
                        label_cell = ws.cell(row=r, column=col_index)
                        value_cell = ws.cell(row=r, column=col_index + 1)

                        label = label_cell.value
                        value = value_cell.value

                        if label in raw_strings and isinstance(value, (int, float)):
                            output_label = rename_map.get(label, label)

                            if threshold is not None and value > threshold:
                                value = value / 1000.0

                            table_values[output_label].append(value)

                    table_data = {}
                    for key, values in table_values.items():
                        if key in ["Table ID", "Sheet_Source"]:
                            table_data[key] = values[0]
                        else:
                            table_data[key] = float(np.mean(values)) if values else None

                    result_data.append(table_data)

    return result_data


def group_data_by_tag(data):
    """
    Groups extracted table entries by time tag extracted from Table ID or Sheet Source.
    """
    w_groups = defaultdict(list)

    for entry in data:
        table_id = entry.get("Table ID", "")
        sheet_source = entry.get("Sheet_Source", "Unknown")
        found_tag = None

        parts = re.split(r'[_\s\-]+', table_id)

        for part in parts:
            if re.fullmatch(r"T\d+", part) or re.fullmatch(r"\d+W", part) or re.fullmatch(r"\d+\+\d+W", part):
                found_tag = part
                break

        if not found_tag:
            found_tag = sheet_source

        w_groups[found_tag or "Unknown"].append(entry)

    return w_groups


def sort_tags(tag):
    """Helper to sort time tags numerically."""
    nums = re.findall(r'\d+', str(tag))
    if nums:
        return int(nums[0])
    return float('inf')
