import os
from collections import Counter
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from . import APP_NAME, __version__
from .parser import find_table_data, normalize_week, parse_table_id, sort_tags, METRIC_COLUMNS, META_KEYS
from .mapping import load_mapping_file, match_key

# Visual Representation Modes
MODE_PASTEL_ACCENTS = "Color genotype cells"
MODE_COLUMNS_ONLY = "No colors"
MODE_SOFT_ROW_FILLS = "Color whole rows"

VISUAL_MODES = [MODE_PASTEL_ACCENTS, MODE_COLUMNS_ONLY, MODE_SOFT_ROW_FILLS]

UNMAPPED = "Unmapped"
LEAD_COLUMNS = ["Table ID", "Animal ID", "Timepoint", "Week", "Genotype", "Sex", "Group", "Runs", "Source Sheet"]

# High-contrast, clean pastel palettes (Hex, Font Color Hex)
MALE_PASTELS = [
    ("EBF3FA", "1B365D"),  # Soft Sky Blue
    ("E3F2FD", "0D47A1"),  # Light Blue
    ("E0F7FA", "006064"),  # Cyan Tint
    ("F0F4C3", "33691E"),  # Soft Lime
]

FEMALE_PASTELS = [
    ("FDE8E8", "7A1C1C"),  # Soft Rose
    ("FCE4EC", "880E4F"),  # Light Pink
    ("FFF3E0", "E65100"),  # Soft Peach
    ("F3E5F5", "4A148C"),  # Soft Lavender
]

NEUTRAL_PASTELS = [
    ("F5F5F5", "212121"),  # Light Gray
    ("FFFDE7", "F57F17"),  # Light Cream
]

MALE_LABELS = {"m", "male", "macho", "masculino"}
FEMALE_LABELS = {"f", "female", "fêmea", "femea", "fémea", "feminino"}

# Common Excel styles
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFFFF")
DATA_FONT = Font(name="Calibri", size=11)
_SIDE = Side(border_style="thin", color="D3D3D3")
CELL_BORDER = Border(left=_SIDE, right=_SIDE, top=_SIDE, bottom=_SIDE)
DATA_HEADER_FILL = PatternFill(start_color="FF2C3E50", end_color="FF2C3E50", fill_type="solid")
SUMMARY_HEADER_FILL = PatternFill(start_color="FF1A5276", end_color="FF1A5276", fill_type="solid")
LOG_HEADER_FILL = PatternFill(start_color="FF7B241C", end_color="FF7B241C", fill_type="solid")


def _get_pastel_fill(genotype, sex, genotype_index):
    """Generates soft pastel PatternFill and Font for group visualization."""
    g = str(genotype).strip().lower()
    s = str(sex).strip().lower()

    if g not in genotype_index:
        return None, None

    if s in MALE_LABELS:
        palette = MALE_PASTELS
    elif s in FEMALE_LABELS:
        palette = FEMALE_PASTELS
    else:
        palette = NEUTRAL_PASTELS

    bg_hex, font_hex = palette[genotype_index[g] % len(palette)]
    return PatternFill(start_color="FF" + bg_hex, end_color="FF" + bg_hex, fill_type="solid"), Font(color="FF" + font_hex, bold=True)


def _short_list(items, limit=30):
    shown = ", ".join(items[:limit])
    return shown + (f", ... (+{len(items) - limit} more)" if len(items) > limit else "")


def build_records(data, mapping_dict=None, threshold=None):
    """
    Turns parsed tables into flat records with metadata, and collects data-quality warnings.

    Returns (records, warnings) where records is a list of dicts with LEAD_COLUMNS + metric columns.
    """
    mapping_dict = mapping_dict or {}
    records, warnings = [], []
    used_keys = set()
    unmapped = []

    for entry in data:
        table_id = entry.get("Table ID", "")
        sheet = entry.get("Sheet_Source", "Unknown")
        animal_id, timepoint, week = parse_table_id(table_id, sheet)

        genotype = sex = None
        if mapping_dict:
            key = match_key(table_id, mapping_dict)
            if key is None:
                unmapped.append(table_id)
            else:
                used_keys.add(key)
                genotype, sex = (v or None for v in mapping_dict[key])
        genotype = genotype or UNMAPPED
        sex = sex or UNMAPPED

        record = {
            "Table ID": table_id,
            "Animal ID": animal_id,
            "Timepoint": timepoint,
            "Week": week,
            "Genotype": genotype,
            "Sex": sex,
            "Group": UNMAPPED if genotype == sex == UNMAPPED else f"{genotype} {sex}",
            "Runs": entry.get("Runs", 0),
            "Source Sheet": sheet,
        }
        for k, v in entry.items():
            if k not in META_KEYS:
                record[k] = v
        records.append(record)

        if not table_id:
            warnings.append(f"Sheet '{sheet}': a table has an empty Table ID.")
        if record["Runs"] == 0:
            warnings.append(f"'{table_id}' (sheet '{sheet}'): no metric values found.")
        else:
            missing = [m for m in METRIC_COLUMNS if m not in entry]
            if missing:
                warnings.append(f"'{table_id}' (sheet '{sheet}'): missing {', '.join(missing)}")
        if entry.get("Converted"):
            warnings.append(
                f"'{table_id}' (sheet '{sheet}'): {entry['Converted']} value(s) larger than {threshold:g} "
                f"were divided by 1000 (decimal-separator fix)."
            )
        sheet_week = normalize_week(sheet)
        if week and sheet_week and week != sheet_week:
            warnings.append(f"'{table_id}' is in sheet '{sheet}' but its Table ID says week {week}.")

    for (table_id, sheet), n in Counter((r["Table ID"], r["Source Sheet"]) for r in records).items():
        if n > 1:
            warnings.append(
                f"'{table_id}' appears {n} times in sheet '{sheet}'; the summary averages these tables."
            )
    if unmapped:
        warnings.append(f"{len(unmapped)} table(s) have no match in the mapping file: {_short_list(unmapped)}")
    unused = [k for k in mapping_dict if k not in used_keys]
    if unused:
        warnings.append(f"{len(unused)} mapping ID(s) matched no table: {_short_list(unused)}")

    return records, warnings


def _metric_cols(df):
    """Metric columns present in df, in canonical order."""
    return [m for m in METRIC_COLUMNS if m in df.columns]


def build_summary(df):
    """One row per animal per timepoint, with the mean of every metric over its tables."""
    metrics = _metric_cols(df)
    keys = ["Animal ID", "Timepoint"]
    agg = {c: "first" for c in ["Table ID", "Week", "Genotype", "Sex", "Group"]}
    agg.update({m: "mean" for m in metrics})
    summary = df.groupby(keys, sort=False, as_index=False).agg(agg)
    summary.insert(2, "Tables", df.groupby(keys, sort=False).size().values)
    return summary[["Animal ID", "Timepoint", "Tables", "Table ID", "Week", "Genotype", "Sex", "Group"] + metrics]


def build_group_stats(summary):
    """N, mean and SEM of every metric per Timepoint x Genotype x Sex, using animals as the unit."""
    metrics = _metric_cols(summary)
    grouped = summary.groupby(["Timepoint", "Genotype", "Sex"], sort=False)
    stats = grouped.size().rename("N animals").to_frame()
    for m in metrics:
        stats[f"{m} Mean"] = grouped[m].mean()
        stats[f"{m} SEM"] = grouped[m].sem()
    return stats.reset_index()


def _write_df(ws, df, header_fill, min_width=12):
    """Writes a DataFrame with styled header, borders, number formats and fitted column widths."""
    ws.append(list(df.columns))
    for row in df.itertuples(index=False, name=None):
        ws.append([None if pd.isna(v) else v for v in row])  # NaN is not valid in xlsx

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = HEADER_FONT
        cell.border = CELL_BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = DATA_FONT
            cell.border = CELL_BORDER
            if isinstance(cell.value, float):
                cell.number_format = "0.000"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    for col in ws.iter_cols():
        max_len = max(len(str(c.value if c.value is not None else "")) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(max_len + 3, min_width), 60)
    ws.freeze_panes = "B2"


def _apply_group_fills(ws, columns, genotype_index, visual_mode):
    """Colors Genotype/Group cells (or whole rows) according to genotype and sex."""
    if visual_mode == MODE_COLUMNS_ONLY or not genotype_index:
        return
    geno_i, sex_i, group_i = (columns.index(c) for c in ("Genotype", "Sex", "Group"))

    for row in ws.iter_rows(min_row=2):
        geno_val, sex_val = row[geno_i].value, row[sex_i].value
        if geno_val == UNMAPPED:
            continue
        fill, font = _get_pastel_fill(geno_val, sex_val, genotype_index)
        if not fill:
            continue
        if visual_mode == MODE_PASTEL_ACCENTS:
            for i in (geno_i, group_i):
                row[i].fill = fill
                row[i].font = font
        elif visual_mode == MODE_SOFT_ROW_FILLS:
            for c in row:
                c.fill = fill


def process_and_export_file(
    input_file,
    mapping_file=None,
    output_dir=None,
    threshold=500.0,
    visual_mode=MODE_PASTEL_ACCENTS,
    generate_summary=True,
    progress_callback=None
):
    """
    Main pipeline function to extract raw Mouseframe data, map metadata,
    apply visual grouping, generate summary statistics per animal/timepoint, and export Excel.

    Returns (output_file, warnings).
    """
    def progress(msg, pct):
        if progress_callback:
            progress_callback(msg, pct)

    if not os.path.isfile(input_file):
        raise FileNotFoundError(f"Input file not found: {input_file}")
    if output_dir and not os.path.isdir(output_dir):
        raise NotADirectoryError(f"Output directory does not exist: {output_dir}")

    progress("Extracting table data from Excel workbook...", 10)
    data = find_table_data(input_file, threshold=threshold)
    if not data:
        raise ValueError("No 'Table ID' tables found in the input file.")

    progress(f"Extracted {len(data)} tables. Loading mapping file...", 30)
    mapping_dict, genotype_index = load_mapping_file(mapping_file)

    records, warnings = build_records(data, mapping_dict, threshold)
    df_all = pd.DataFrame(records)
    metrics = _metric_cols(df_all)
    if not metrics:
        raise ValueError("Tables were found, but none contain the expected 'average in cm' metric rows.")
    df_all = df_all[LEAD_COLUMNS + metrics]

    now = datetime.now()
    timestamp = now.strftime("%Y-%m-%d_%H-%M-%S")
    base_dir = output_dir or os.path.dirname(os.path.abspath(input_file))
    output_file = os.path.normpath(os.path.join(base_dir, f"MFrame_clean_output_{timestamp}.xlsx"))

    wb = Workbook()
    wb.remove(wb.active)

    if generate_summary:
        progress("Generating Summary_Per_Animal sheet...", 40)
        summary = build_summary(df_all)
        ws = wb.create_sheet("Summary_Per_Animal")
        _write_df(ws, summary, SUMMARY_HEADER_FILL, min_width=14)
        _apply_group_fills(ws, list(summary.columns), genotype_index, visual_mode)

        if mapping_dict:
            stats = build_group_stats(summary)
            ws = wb.create_sheet("Group_Statistics")
            _write_df(ws, stats, SUMMARY_HEADER_FILL, min_width=14)

    tags = sorted(df_all["Timepoint"].unique(), key=sort_tags)
    for i, tag in enumerate(tags, start=1):
        progress(f"Writing timepoint sheet '{tag}' ({i}/{len(tags)})...", 45 + int(45 * i / len(tags)))
        df = df_all[df_all["Timepoint"] == tag].drop(columns="Timepoint")
        ws = wb.create_sheet(tag)
        _write_df(ws, df, DATA_HEADER_FILL)
        _apply_group_fills(ws, list(df.columns), genotype_index, visual_mode)

    if warnings:
        ws = wb.create_sheet("Processing_Log")
        _write_df(ws, pd.DataFrame({"Warning": warnings}), LOG_HEADER_FILL)

    # Provenance: which inputs and settings produced this workbook
    info = [
        ("Created by", f"{APP_NAME} {__version__}"),
        ("Created on", now.strftime("%Y-%m-%d %H:%M:%S")),
        ("Input file", os.path.abspath(input_file)),
        ("Mapping file", os.path.abspath(mapping_file) if mapping_file else "(none)"),
        ("Decimal-fix threshold", f"{threshold:g}"),
        ("Color style", visual_mode),
        ("Summary sheets", "Yes" if generate_summary else "No"),
        ("Tables", len(records)),
        ("Animals", df_all["Animal ID"].nunique()),
        ("Timepoints", len(tags)),
        ("Warnings", len(warnings)),
    ]
    _write_df(wb.create_sheet("Info"), pd.DataFrame(info, columns=["Setting", "Value"]), SUMMARY_HEADER_FILL)

    progress("Saving workbook to disk...", 95)
    # Save to a temp file first so a crash or closed app never leaves a half-written workbook
    tmp_file = output_file + ".part"
    try:
        wb.save(tmp_file)
        os.replace(tmp_file, output_file)
    finally:
        if os.path.exists(tmp_file):
            os.remove(tmp_file)

    progress(f"Done. {len(warnings)} warning(s).", 100)
    return output_file, warnings
