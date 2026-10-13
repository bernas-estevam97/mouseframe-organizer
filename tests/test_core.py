"""Core pipeline checks. Run with:  python -m unittest discover tests"""
import os
import sys
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from core.parser import find_table_data, parse_table_id, normalize_week, sort_tags
from core.mapping import load_mapping_file, match_table_id
from core.excel_writer import build_records, process_and_export_file

SL_LF = "Stride length left front average in cm:"
OV_R = "Overlap Right average in cm:"
SW_FL = "Stride Width Front(L) average in cm:"
SW_FR = "Stride Width Front(R) average in cm:"


def _table(ws, col, table_id, rows):
    """Writes a MouseFrame-style table: 'Table ID' in row 1, (label, value) pairs below."""
    ws.cell(1, col, f"Table ID: {table_id}")
    for i, (label, value) in enumerate(rows, start=2):
        ws.cell(i, col, label)
        if value is not None:
            ws.cell(i, col + 1, value)


class CoreTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = self.tmp.name
        wb = Workbook()
        ws = wb.active
        ws.title = "4W"
        # two runs separated by BREAK -> averaged, Runs=2
        _table(ws, 1, "2199_1R_T1_4W", [(SL_LF, 8.0), (SW_FL, 1.0), ("---BREAK---", None), (SL_LF, 10.0), (SW_FL, 3.0)])
        # comma decimal string, value above threshold, missing Overlap Right
        _table(ws, 8, "2200_1L_T1_4W", [(SL_LF, "7,5"), (SW_FR, 1200)])
        # duplicate of the first table in the same sheet
        _table(ws, 15, "2199_1R_T1_4W", [(SL_LF, 6.0), (SW_FL, 2.0)])
        ws2 = wb.create_sheet("56W+1")
        _table(ws2, 1, "3220_T14_56+1W", [(SL_LF, 9.0), (OV_R, 0.5)])
        # a second table stacked below in the same column must not leak into the first
        ws2.cell(10, 1, "Table ID: 3223_T14_56+1W")
        ws2.cell(11, 1, SL_LF)
        ws2.cell(11, 2, 5.0)
        self.xlsx = os.path.join(self.dir, "raw.xlsx")
        wb.save(self.xlsx)

        self.map_csv = os.path.join(self.dir, "map.csv")
        with open(self.map_csv, "w", encoding="utf-8") as f:
            f.write("ID,Genotype,Sex\n2199,WT,M\n2199_1R,SCA2,F\n2200,WT,F\n3220,SCA2,M\n3223,SCA2,M\n9999,WT,M\n")

    def tearDown(self):
        self.tmp.cleanup()

    def test_parse_tables(self):
        data = {(d["Sheet_Source"], d["Table ID"], i): d for i, d in enumerate(find_table_data(self.xlsx))}
        self.assertEqual(len(data), 5)
        first = next(d for k, d in data.items() if k[1] == "2199_1R_T1_4W")
        self.assertEqual(first["Runs"], 2)
        self.assertAlmostEqual(first[SL_LF], 9.0)
        self.assertAlmostEqual(first["Stride Width Front average in cm:"], 2.0)
        second = next(d for k, d in data.items() if k[1] == "2200_1L_T1_4W")
        self.assertAlmostEqual(second[SL_LF], 7.5)
        self.assertAlmostEqual(second["Stride Width Front average in cm:"], 1.2)
        self.assertEqual(second["Converted"], 1)
        stacked = next(d for k, d in data.items() if k[1] == "3220_T14_56+1W")
        self.assertAlmostEqual(stacked[SL_LF], 9.0)

    def test_rejects_xls(self):
        with self.assertRaises(ValueError):
            find_table_data(os.path.join(self.dir, "old.xls"))

    def test_table_id_parts(self):
        self.assertEqual(parse_table_id("2199_1R_T1_4W"), ("2199_1R", "T1", "4W"))
        self.assertEqual(parse_table_id("3220_T17_68W+1"), ("3220", "T17", "68W+1"))
        self.assertEqual(parse_table_id("3220_T14_56+1W"), ("3220", "T14", "56W+1"))
        self.assertEqual(parse_table_id("3220_T14_56W+3"), ("3220", "T14", "56W+3"))  # +N weeks share the session
        self.assertEqual(parse_table_id("3220_56W", "Sheet9"), ("3220", "56W", "56W"))
        self.assertEqual(parse_table_id("mouse7", "Day 3"), ("mouse7", "Day 3", ""))
        self.assertEqual(normalize_week("56W+2"), "56W+2")
        self.assertIsNone(normalize_week("1R"))
        self.assertEqual(sorted(["T10", "T2", "T1", "Unknown"], key=sort_tags), ["T1", "T2", "T10", "Unknown"])

    def test_mapping_match(self):
        mapping, geno_index = load_mapping_file(self.map_csv)
        self.assertEqual(match_table_id("2199_1R_T1_4W", mapping), ("SCA2", "F"))  # longest key wins
        self.assertEqual(match_table_id("2199_1L_T1_4W", mapping), ("WT", "M"))
        self.assertEqual(match_table_id("21990_T1_4W", mapping), (None, None))  # no partial-number match
        self.assertEqual(set(geno_index), {"wt", "sca2"})

    def test_mapping_errors(self):
        with self.assertRaises(FileNotFoundError):
            load_mapping_file(os.path.join(self.dir, "missing.csv"))
        bad = os.path.join(self.dir, "bad.csv")
        with open(bad, "w") as f:
            f.write("ID,Genotype,Sex\n1,WT,M\n1,KO,M\n")
        with self.assertRaises(ValueError):
            load_mapping_file(bad)

    def test_records_and_warnings(self):
        mapping, _ = load_mapping_file(self.map_csv)
        records, warnings = build_records(find_table_data(self.xlsx), mapping, 500.0)
        text = "\n".join(warnings)
        self.assertIn("appears 2 times", text)
        self.assertIn("divided by 1000", text)
        self.assertIn("missing", text)
        self.assertIn("9999", text)  # unused mapping key
        self.assertTrue(all(r["Genotype"] != "Unmapped" for r in records))

    def test_export(self):
        out, warnings = process_and_export_file(self.xlsx, self.map_csv, self.dir)
        wb = load_workbook(out)
        self.assertEqual(wb.sheetnames, ["Summary_Per_Animal", "Group_Statistics", "T1", "T14", "Processing_Log"])
        summary = list(wb["Summary_Per_Animal"].iter_rows(values_only=True))
        header = summary[0]
        row = next(r for r in summary if r[0] == "2199_1R")
        self.assertEqual(row[header.index("Tables")], 2)
        self.assertAlmostEqual(row[header.index(SL_LF)], 7.5)  # mean of 9.0 and 6.0
        self.assertFalse(any(f.endswith(".part") for f in os.listdir(self.dir)))

    def test_export_bad_output_dir(self):
        with self.assertRaises(NotADirectoryError):
            process_and_export_file(self.xlsx, output_dir=os.path.join(self.dir, "nope"))


if __name__ == "__main__":
    unittest.main()
